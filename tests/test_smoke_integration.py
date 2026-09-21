"""End-to-end integration tests: real file -> DeepEngine -> Markdown -> chunks.

Unlike the unit tests (which mock the engines out), these drive the *real*
Deep Engine against genuinely-constructed Office files so the full RAG path
(engine → pipeline → cleaner → chunker) is proven, not just its seams. They
need the real deps (openpyxl / python-docx) that the core install already
ships, so they are part of the default suite, not an optional add-on.
"""

from pathlib import Path

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook

from omnidoc.controller import ConversionController
from omnidoc.core.config import OmniDocConfig


def _cfg() -> OmniDocConfig:
    cfg = OmniDocConfig.default()
    cfg.chunking.enabled = True
    cfg.chunking.strategy = "fixed"
    cfg.chunking.chunk_size = 400
    cfg.chunking.chunk_overlap = 40
    cfg.cleaning.remove_empty_lines = True
    cfg.cleaning.normalize_spaces = True
    return cfg


def _make_xlsx(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Product", "Units", "Revenue"])
    ws.append(["North", "Widget A", 1200, 360000.0])
    ws.append(["North", "Widget B", 980, 294000.0])
    ws.append(["South", "Widget A", 760, 228000.0])
    ws.append(["South", "Widget B", 640, 192000.0])
    ws.append(["Total", "", 3580, 1074000.0])
    wb.save(path)
    return path


def _make_docx(path: Path) -> Path:
    doc = DocxDocument()
    doc.add_heading("Q1 2026 Report", 1)
    doc.add_paragraph("This quarter total revenue reached 1,074,000 CNY.")
    p = doc.add_paragraph()
    p.add_run("North").bold = True
    p.add_run(" region led the total with 654,000 CNY.")
    doc.add_heading("Widget A", level=2)
    doc.add_paragraph("1,960 units sold across both regions.")
    doc.save(path)
    return path


class TestXlsxIntegration:

    def test_real_xlsx_produces_markdown_and_chunks(self, tmp_path: Path):
        xlsx = _make_xlsx(tmp_path / "q1.xlsx")
        result = ConversionController().convert(str(xlsx), _cfg())

        assert result.success
        assert result.status.value == "ok"
        assert result.engine == "deep"
        assert result.fallback_used is False
        assert not result.errors
        assert not result.warnings
        assert "Region" in result.markdown
        assert "Widget B" in result.markdown
        assert len(result.chunks) >= 1
        # chunk metadata is self-describing
        c0 = result.chunks[0]
        assert c0.metadata["chunk_index"] == 0
        assert c0.metadata["chunk_count"] == len(result.chunks)

    def test_to_dict_shape(self, tmp_path: Path):
        xlsx = _make_xlsx(tmp_path / "q1.xlsx")
        result = ConversionController().convert(str(xlsx), _cfg())
        payload = result.to_dict()
        for key in ("source", "engine", "status", "success", "chunk_count", "chunks", "markdown"):
            assert key in payload
        assert payload["chunk_count"] == len(result.chunks)
        assert all("text" in c for c in payload["chunks"])


class TestDocxIntegration:

    def test_real_docx_produces_markdown_and_chunks(self, tmp_path: Path):
        docx = _make_docx(tmp_path / "report.docx")
        result = ConversionController().convert(str(docx), _cfg())

        assert result.success
        assert result.engine == "deep"
        assert result.fallback_used is False
        assert not result.errors
        assert "Q1 2026 Report" in result.markdown
        assert "Widget A" in result.markdown
        assert len(result.chunks) >= 1


class TestBatchIntegration:

    def test_batch_mixed_sources_one_bad_file_does_not_abort(self, tmp_path: Path):
        good = _make_xlsx(tmp_path / "good.xlsx")
        bad = _make_xlsx(tmp_path / "bad.xlsx")
        # truncate the "bad" file to break it
        bad.write_bytes(b"garbage-not-a-zip")
        controller = ConversionController()
        results = controller.convert_batch([str(good), str(bad)], _cfg())

        assert len(results) == 2
        assert results[0].success
        assert results[0].engine == "deep"
        # the corrupt file degrades, never aborts the batch
        assert results[1].success is False or results[1].errors


class TestDeepAllSheetsFail:

    def test_all_sheets_failure_reports_error_status(self, tmp_path: Path):
        # A corrupt .xlsx makes load_sheets raise per-sheet, so build() returns
        # {"sheets": [], "errors": [...]}. DeepEngine must surface ERROR status,
        # not the default OK, on that branch.
        from omnidoc.engines.deep_engine import DeepEngine
        from omnidoc.core.document import ConversionStatus

        bad = tmp_path / "corrupt.xlsx"
        bad.write_bytes(b"garbage-not-a-zip")
        eng = DeepEngine()
        res = eng._result_multi(
            {"sheets": [], "errors": [("Sales", "Invalid file: boom")],
             "stem": "corrupt", "source_name": "corrupt.xlsx", "metadata": {"sheet_count": 0}},
            "md", str(bad), {}, None,
        )
        assert res.status == ConversionStatus.ERROR
        assert not res.success
        assert any("Sales" in e for e in res.errors)


class TestDegradation:

    def test_nonexistent_source_degrades_to_error_result(self, tmp_path: Path):
        # A missing file is caught by the Deep Engine's fallback chain and
        # recorded as a non-fatal error (System Rule #3: never abort, never
        # raise). The pipeline therefore yields NO markdown (empty string —
        # the cleaner/chunker stages short-circuit on an empty result) while
        # the diagnostic is preserved on result.errors, so result.success is
        # False without the whole run ever raising.
        missing = tmp_path / "does_not_exist.xlsx"
        result = ConversionController().convert(str(missing), _cfg())
        assert result.success is False
        assert result.markdown == ""
        assert result.errors
        assert not result.chunks


if __name__ == "__main__":
    import sys

    sys.exit(pytest.main([__file__, "-v"]))
